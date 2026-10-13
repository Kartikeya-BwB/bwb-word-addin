"""Builds the BwB PowerPoint template, its slide library and the library index the add-in reads.

Usage:
  python3 scripts/make-slide-library.py [--no-thumbs] [--preview DIR]

Writes, relative to source/:
  public/templates/bwb-slides.pptx       a 16:9 BwB slide master (theme, seven layouts) and the library slides
  public/templates/bwb-deck.pptx         the same master with a blank title slide and a blank title and content
                                         slide: the starting point for a new BwB deck
  public/assets/slides/index.json        the template, the starter deck (starterDeck) and one entry per slide: id,
                                         title, category, description, tags, slideId, thumb
  public/assets/slides/thumbs/<id>.png   480 px thumbnails, rendered with LibreOffice and pdftoppm (palette PNG)

Positions come from src/brand/slideGrid.ts and colours from src/brand/tokens.ts, read when the script runs, so the
template keeps to the grid the add-in's alignment tools use. Photos come from public/assets/images, the logos from
public/assets/brand and the icons from lucide-static (rendered with cairosvg, or with @resvg/resvg-js when cairosvg is
not installed). Icons go in as SVG with a PNG fallback, so PowerPoint keeps them sharp and recolourable.

Fonts are named, never embedded: Candara and Aptos must be installed for the thumbnails to be faithful.

Every slide has a fixed sldId and a p14:creationId derived from its library id, so the add-in's slide ids
("<sldId>#<creationId>") stay the same from one build to the next. Both files are written with fixed timestamps,
the workbook behind the native chart included, so rebuilding without changes gives the same bytes.
"""
import argparse
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import zlib
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml import parse_xml
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent
PPTX_OUT = ROOT / "public" / "templates" / "bwb-slides.pptx"
DECK_OUT = ROOT / "public" / "templates" / "bwb-deck.pptx"
SLIDES_OUT = ROOT / "public" / "assets" / "slides"
IMAGES = ROOT / "public" / "assets" / "images"
BRAND = ROOT / "public" / "assets" / "brand"
LUCIDE = ROOT / "node_modules" / "lucide-static" / "icons"
FIXED_TIME = (2026, 10, 13, 9, 0, 0)
FIXED_ISO = "2026-10-13T09:00:00Z"
THUMB_WIDTH = 480

NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
P15 = "http://schemas.microsoft.com/office/powerpoint/2012/main"
CREATION_EXT = "{BB962C8B-B14F-4D97-AF65-F5344CB8AC3E}"
GUIDE_EXT = "{27BBF7A9-308A-43DC-89C8-2F10F3537804}"
SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
SLIDENUM_FIELD = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"
BWB_TABLE_STYLE = "{6A0B2E4D-5C3F-4E8B-9C1A-2B7D4F6E8A10}"
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
LAYOUT_NAMES = ["BwB Title", "BwB Section divider", "BwB Title and content", "BwB Two content", "BwB Title only",
                "BwB Blank", "BwB Closing"]
NAVY_LAYOUTS = ("BwB Title", "BwB Section divider", "BwB Closing")
WHITE_LOGO_LAYOUTS = ("BwB Title", "BwB Closing")
CATEGORIES = ("Structure", "Messages", "Data", "Visuals", "People", "Closing")


# ---------------------------------------------------------------------------------------------------------------------
# The grid and the palette, read from the add-in's TypeScript so the two cannot drift apart

class Grid:
    """The BwB slide grid in points (src/brand/slideGrid.ts)."""

    def __init__(self, ts: str):
        num = r"([\d.]+)"

        def find(pattern: str, what: str) -> list:
            m = re.search(pattern, ts)
            if not m:
                sys.exit(f"Could not read {what} from src/brand/slideGrid.ts")
            return [float(v) for v in m.groups()]

        def area(name: str) -> tuple:
            return tuple(find(rf"export const {name}: Area = \{{ left: {num}, top: {num}, width: {num}, height: {num} \}}",
                              name))

        self.slide, self.title, self.content = area("SLIDE"), area("TITLE_AREA"), area("CONTENT_AREA")
        self.source, self.footer = area("SOURCE_AREA"), area("FOOTER_AREA")
        left, right, top, bottom = find(rf"export const MARGINS = \{{ left: {num}, right: {num}, top: {num}, "
                                        rf"bottom: {num} \}}", "MARGINS")
        self.margin_left, self.margin_right, self.margin_top, self.margin_bottom = left, right, top, bottom
        self.strip = find(rf"export const STRIP_HEIGHT = {num}", "STRIP_HEIGHT")[0]
        self.columns = int(find(rf"export const COLUMNS = {num}", "COLUMNS")[0])
        self.gutter = find(rf"export const GUTTER = {num}", "GUTTER")[0]
        self.width, self.height = self.slide[2], self.slide[3]
        safe = self.width - left - right
        self.col = (safe - (self.columns - 1) * self.gutter) / self.columns

    def x(self, column: int) -> float:
        """Left edge of a column, counting from 0."""
        return self.margin_left + column * (self.col + self.gutter)

    def w(self, columns: int) -> float:
        """Width of a span of whole columns and the gutters between them."""
        return columns * self.col + (columns - 1) * self.gutter


def read_colours() -> dict:
    ts = (ROOT / "src" / "brand" / "tokens.ts").read_text(encoding="utf-8")
    block = re.search(r"export const COLOURS = \{(.*?)\} as const", ts, re.S)
    if not block:
        sys.exit("Could not read COLOURS from src/brand/tokens.ts")
    return {k: v.upper() for k, v in re.findall(r"(\w+): '([0-9A-Fa-f]{6})'", block.group(1))}


G = Grid((ROOT / "src" / "brand" / "slideGrid.ts").read_text(encoding="utf-8"))
C = read_colours()
NAVY, GREEN, LIME, CHARCOAL = C["navy"], C["green"], C["lime"], C["charcoal"]
GREEN50, GREY, OFFWHITE, WHITE, BLACK, GREY20 = (C["green50"], C["lightGrey"], C["offWhite"], C["white"], C["black"],
                                                 C["charcoal20"])
HEAD, BODY = "Candara", "Aptos"
CT = G.content[1]                    # content top
CB = G.content[1] + G.content[3]     # content bottom
CH = G.content[3]
RIGHT = G.width - G.margin_right


# ---------------------------------------------------------------------------------------------------------------------
# DrawingML helpers. Everything is in points; emu() converts.

def emu(pt: float) -> int:
    return int(round(pt * 12700))


def esc(text: str) -> str:
    return escape(text, {'"': "&quot;"})


def solid(colour: str) -> str:
    return f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill>'


def xfrm(box: tuple, tag: str = "a:xfrm", flip_v: bool = False) -> str:
    x, y, w, h = box
    flip = ' flipV="1"' if flip_v else ""
    return f'<{tag}{flip}><a:off x="{emu(x)}" y="{emu(y)}"/><a:ext cx="{emu(w)}" cy="{emu(h)}"/></{tag}>'


def rpr(size=None, colour=None, font=None, bold=None, italic=None, tag="a:rPr") -> str:
    attrs = ['lang="en-GB"']
    if size is not None:
        attrs.append(f'sz="{int(round(size * 100))}"')
    if bold is not None:
        attrs.append(f'b="{1 if bold else 0}"')
    if italic is not None:
        attrs.append(f'i="{1 if italic else 0}"')
    inner = solid(colour) if colour else ""
    if font:
        inner += f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/>'
    return f'<{tag} {" ".join(attrs)}>{inner}</{tag}>'


def runs(text: str, size=None, colour=None, font=None, bold=None, italic=None) -> str:
    """Runs for one paragraph. **Bold** marks bold words and a vertical tab is a line break."""
    out = []
    for li, line in enumerate(text.split("\v")):
        if li:
            out.append(f'<a:br>{rpr(size, colour, font, bold, italic)}</a:br>')
        for i, part in enumerate(line.split("**")):
            if part:
                strong = True if i % 2 else bold
                out.append(f'<a:r>{rpr(size, colour, font, strong, italic)}<a:t>{esc(part)}</a:t></a:r>')
    return "".join(out)


# bullet: colour, character, font, size in per cent, left margin and hanging indent in points. The glyphs, fonts and
# colours are the ones the BwB Word template uses for its lists.
BULLETS = {
    "square": (GREEN, "\u25A0", "Arial", 80, 16, 16),
    "circle": (NAVY, "\u25CF", "Arial", 70, 32, 14),
    "dash": (CHARCOAL, "\u2013", "Arial", 100, 46, 14),
}


def bullet_xml(bullet: str) -> tuple:
    """The margin attributes and bullet elements of a paragraph or list level."""
    if not bullet:
        return 'marL="0" indent="0"', "<a:buNone/>"
    colour, char, font, pct, marl, hang = BULLETS[bullet]
    return (f'marL="{emu(marl)}" indent="{emu(-hang)}"',
            f'<a:buClr><a:srgbClr val="{colour}"/></a:buClr><a:buSzPct val="{pct * 1000}"/>'
            f'<a:buFont typeface="{font}" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
            f'<a:buChar char="{char}"/>')


def P(text: str = "", size: float = 14, colour: str = BLACK, font: str = BODY, bold: bool = False,
      italic: bool = False, align: str = "l", after: float = 0, before: float = 0, line: float = 0,
      bullet: str = "") -> str:
    """One paragraph with explicit formatting, for text boxes and shapes."""
    attrs, bu = bullet_xml(bullet)
    spacing = f'<a:lnSpc><a:spcPct val="{int(line * 1000)}"/></a:lnSpc>' if line else ""
    spacing += f'<a:spcBef><a:spcPts val="{int(before * 100)}"/></a:spcBef>'
    spacing += f'<a:spcAft><a:spcPts val="{int(after * 100)}"/></a:spcAft>'
    ppr = f'<a:pPr {attrs} algn="{align}">{spacing}{bu}</a:pPr>'
    end = rpr(size, colour, font, tag="a:endParaRPr")
    return f"<a:p>{ppr}{runs(text, size, colour, font, bold or None, italic or None)}{end}</a:p>"


def PH(text: str, level: int = 0) -> str:
    """One paragraph for a placeholder: the formatting comes from the layout and the master."""
    ppr = f'<a:pPr lvl="{level}"/>' if level else ""
    return f'<a:p>{ppr}{runs(text)}<a:endParaRPr lang="en-GB"/></a:p>'


def body_pr(anchor: str = "t", inset: tuple = (0, 0, 0, 0), vert: str = "", autofit: str = "<a:noAutofit/>") -> str:
    l, t, r, b = inset
    rotate = f' vert="{vert}"' if vert else ""
    return (f'<a:bodyPr{rotate} wrap="square" lIns="{emu(l)}" tIns="{emu(t)}" rIns="{emu(r)}" bIns="{emu(b)}" '
            f'rtlCol="0" anchor="{anchor}" anchorCtr="0">{autofit}</a:bodyPr>')


def lvl_ppr(level: int, size: float, colour: str, font: str = "+mn-lt", bold: bool = False, italic: bool = False,
            align: str = "l", line: float = 0, after: float = 0, bullet: str = "") -> str:
    """A list style level (lvlNpPr) for layout placeholders and the master's text styles."""
    attrs, bu = bullet_xml(bullet)
    spacing = f'<a:lnSpc><a:spcPct val="{int(line * 1000)}"/></a:lnSpc>' if line else ""
    spacing += f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="{int(after * 100)}"/></a:spcAft>'
    style = ' b="1"' if bold else ' b="0"'
    style += ' i="1"' if italic else ""
    ea, cs = ("+mj-ea", "+mj-cs") if font.startswith("+mj") else ("+mn-ea", "+mn-cs") if font.startswith("+") else (
        "+mn-ea", font)
    rp = (f'<a:defRPr lang="en-GB" sz="{int(size * 100)}"{style} kern="1200">{solid(colour)}'
          f'<a:latin typeface="{font}"/><a:ea typeface="{ea}"/><a:cs typeface="{cs}"/></a:defRPr>')
    return (f'<a:lvl{level}pPr {attrs} algn="{align}" defTabSz="914400" rtl="0" eaLnBrk="1" latinLnBrk="0" '
            f'hangingPunct="1">{spacing}{bu}{rp}</a:lvl{level}pPr>')


class Canvas:
    """Collects the shapes of one slide, layout or master and numbers them from 2."""

    def __init__(self):
        self.parts = []
        self.last = 1

    def _id(self) -> int:
        self.last += 1
        return self.last

    def shape(self, name: str, box: tuple, *, geom: str = "rect", adj: dict = None, fill: str = None,
              line: str = None, line_w: float = 0.75, text: str = None, anchor: str = "t",
              inset: tuple = (0, 0, 0, 0), txbox: bool = False, vert: str = "") -> int:
        sid = self._id()
        tb = ' txBox="1"' if txbox else ""
        guides = "".join(f'<a:gd name="{k}" fmla="val {v}"/>' for k, v in (adj or {}).items())
        ln = f'<a:ln w="{emu(line_w)}">{solid(line)}</a:ln>' if line else "<a:ln><a:noFill/></a:ln>"
        sp_pr = (f'<p:spPr>{xfrm(box)}<a:prstGeom prst="{geom}"><a:avLst>{guides}</a:avLst></a:prstGeom>'
                 f'{solid(fill) if fill else "<a:noFill/>"}{ln}</p:spPr>')
        tx = f"<p:txBody>{body_pr(anchor, inset, vert)}<a:lstStyle/>{text}</p:txBody>" if text is not None else ""
        self.parts.append(f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{esc(name)}"/><p:cNvSpPr{tb}/><p:nvPr/>'
                          f"</p:nvSpPr>{sp_pr}{tx}</p:sp>")
        return sid

    def text(self, name: str, box: tuple, paras: str, **kw) -> int:
        return self.shape(name, box, text=paras, txbox=True, **kw)

    def line(self, name: str, x1: float, y1: float, x2: float, y2: float, colour: str, width: float = 0.75,
             arrow: bool = False) -> int:
        sid = self._id()
        box = (min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
        head = '<a:tailEnd type="triangle" w="med" len="med"/>' if arrow else ""
        self.parts.append(
            f'<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="{sid}" name="{esc(name)}"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
            f'<p:spPr>{xfrm(box, flip_v=y2 < y1)}<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
            f'<a:ln w="{emu(width)}">{solid(colour)}{head}</a:ln></p:spPr></p:cxnSp>')
        return sid

    def picture(self, name: str, rid: str, box: tuple, descr: str, svg_rid: str = None) -> int:
        sid = self._id()
        ext = ""
        if svg_rid:
            ext = (f'<a:extLst><a:ext uri="{SVG_EXT}"><asvg:svgBlip '
                   f'xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="{svg_rid}"/>'
                   "</a:ext></a:extLst>")
        self.parts.append(
            f'<p:pic><p:nvPicPr><p:cNvPr id="{sid}" name="{esc(name)}" descr="{esc(descr)}"/><p:cNvPicPr>'
            '<a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
            f'<p:blipFill><a:blip r:embed="{rid}">{ext}</a:blip><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            f'<p:spPr>{xfrm(box)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
        return sid

    def placeholder(self, name: str, ph: str, paras: str = "", box: tuple = None, bpr: str = "<a:bodyPr/>",
                    lst: str = "") -> int:
        sid = self._id()
        sp_pr = f"<p:spPr>{xfrm(box)}</p:spPr>" if box else "<p:spPr/>"
        self.parts.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{esc(name)}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f"<p:nvPr><p:ph {ph}/></p:nvPr></p:nvSpPr>{sp_pr}<p:txBody>{bpr}<a:lstStyle>{lst}</a:lstStyle>"
            f"{paras or PH('')}</p:txBody></p:sp>")
        return sid

    def slide_number(self, ph: str, number: str = "\u2039#\u203A", box: tuple = None, bpr: str = "<a:bodyPr/>",
                     lst: str = "") -> int:
        field = (f'<a:p><a:fld id="{SLIDENUM_FIELD}" type="slidenum"><a:rPr lang="en-GB"/><a:t>{number}</a:t></a:fld>'
                 '<a:endParaRPr lang="en-GB"/></a:p>')
        return self.placeholder("Slide number", ph, field, box, bpr, lst)

    def table(self, name: str, x: float, y: float, widths: list, rows: list) -> int:
        """rows: list of (height, cells); a cell is the <a:tc> XML from tc()."""
        sid = self._id()
        grid = "".join(f'<a:gridCol w="{emu(w)}"/>' for w in widths)
        body = "".join(f'<a:tr h="{emu(h)}">{"".join(cells)}</a:tr>' for h, cells in rows)
        box = (x, y, sum(widths), sum(h for h, _ in rows))
        self.parts.append(
            f'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="{sid}" name="{esc(name)}"/><p:cNvGraphicFramePr>'
            '<a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
            f'{xfrm(box, "p:xfrm")}<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
            f'<a:tbl><a:tblPr firstRow="1" bandRow="0"><a:tableStyleId>{NO_STYLE_TABLE}</a:tableStyleId></a:tblPr>'
            f"<a:tblGrid>{grid}</a:tblGrid>{body}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>")
        return sid

    def chart(self, name: str, rid: str, box: tuple, descr: str) -> int:
        sid = self._id()
        self.parts.append(
            f'<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="{sid}" name="{esc(name)}" descr="{esc(descr)}"/>'
            '<p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
            f'{xfrm(box, "p:xfrm")}<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            f'<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" r:id="{rid}"/></a:graphicData>'
            "</a:graphic></p:graphicFrame>")
        return sid

    def strip(self):
        """The green, charcoal and lime strip along the bottom edge, as in the BwB documents."""
        third = G.width / 3
        for i, (label, colour) in enumerate((("green", GREEN), ("charcoal", CHARCOAL), ("lime", LIME))):
            self.shape(f"Colour strip {label}", (i * third, G.height - G.strip, third, G.strip), fill=colour)

    def xml(self) -> str:
        return "".join(self.parts)


def tc(paras: str, fill: str = None, lines: dict = None, margins: tuple = (6, 4, 6, 4), anchor: str = "ctr",
       span: int = 0, hmerge: bool = False, rowspan: int = 0, vmerge: bool = False) -> str:
    """One table cell. lines maps L, R, T and B to (colour, width in points)."""
    attrs = (f' gridSpan="{span}"' if span else "") + (f' rowSpan="{rowspan}"' if rowspan else "")
    attrs += (' hMerge="1"' if hmerge else "") + (' vMerge="1"' if vmerge else "")
    borders = ""
    for side in "LRTB":
        spec = (lines or {}).get(side)
        if spec:
            borders += (f'<a:ln{side} w="{emu(spec[1])}" cap="flat" cmpd="sng" algn="ctr">{solid(spec[0])}'
                        f'<a:prstDash val="solid"/></a:ln{side}>')
        else:
            borders += f'<a:ln{side} w="0"><a:noFill/></a:ln{side}>'
    l, t, r, b = margins
    return (f'<a:tc{attrs}><a:txBody><a:bodyPr/><a:lstStyle/>{paras}</a:txBody>'
            f'<a:tcPr marL="{emu(l)}" marR="{emu(r)}" marT="{emu(t)}" marB="{emu(b)}" anchor="{anchor}">'
            f'{borders}{solid(fill) if fill else "<a:noFill/>"}</a:tcPr></a:tc>')


def creation_id(key: str) -> int:
    """A stable p14:creationId for a slide, layout or master, derived from its name."""
    return (zlib.crc32(f"bwb-slides/{key}".encode("utf-8")) & 0x7FFFFFFF) or 1


def creation_ext(key: str) -> str:
    return (f'<p:extLst><p:ext uri="{CREATION_EXT}"><p14:creationId xmlns:p14="{P14}" val="{creation_id(key)}"/>'
            "</p:ext></p:extLst>")


def sp_tree(content: str) -> str:
    return ('<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr>'
            '<a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm>'
            f"</p:grpSpPr>{content}</p:spTree>")


# ---------------------------------------------------------------------------------------------------------------------
# Brand assets: logos trimmed to their visible edges, cropped photos and recoloured Lucide icons

def trimmed_png(path: Path) -> tuple:
    with Image.open(path) as im:
        im = im.convert("RGBA")
        im = im.crop(im.getchannel("A").getbbox())
        buf = io.BytesIO()
        im.save(buf, "PNG", optimize=True)
        return buf.getvalue(), im.size[0] / im.size[1]


def photo(image_id: str, aspect: float, cx: float, cy: float, scale: float, width_px: int) -> bytes:
    """Crops a library photo to the frame's aspect ratio around (cx, cy), given as fractions of the photo.

    scale is the crop's width as a fraction of the photo's width (capped by the photo's height)."""
    with Image.open(IMAGES / f"{image_id}.jpg") as im:
        w, h = im.size
        cw = min(w * scale, h * aspect, w)
        ch = cw / aspect
        left = min(max(cx * w - cw / 2, 0), w - cw)
        top = min(max(cy * h - ch / 2, 0), h - ch)
        crop = im.crop((round(left), round(top), round(left + cw), round(top + ch)))
        crop = crop.convert("RGB").resize((width_px, round(width_px / aspect)), Image.LANCZOS)
        buf = io.BytesIO()
        crop.save(buf, "JPEG", quality=82, optimize=True)
        return buf.getvalue()


def lucide_svg(name: str, colour: str) -> bytes:
    svg = (LUCIDE / f"{name}.svg").read_text(encoding="utf-8")
    svg = re.sub(r'\s*class="[^"]*"', "", svg).replace("currentColor", f"#{colour}")
    return svg.encode("utf-8")


def svg_to_png(svg: bytes, width: int) -> bytes:
    try:
        import cairosvg
        return cairosvg.svg2png(bytestring=svg, output_width=width)
    except ImportError:
        script = ("const {Resvg}=require('@resvg/resvg-js');let s='';process.stdin.on('data',d=>s+=d);"
                  "process.stdin.on('end',()=>{const r=new Resvg(s,{fitTo:{mode:'width',value:+process.argv[1]}});"
                  "process.stdout.write(r.render().asPng())})")
        return subprocess.run(["node", "-e", script, str(width)], input=svg, cwd=ROOT, capture_output=True,
                              check=True).stdout


class Assets:
    """Adds pictures, icons and charts to one slide part and hands back their relationship ids."""

    def __init__(self, part, package):
        self.part, self.package = part, package

    def image(self, data: bytes) -> str:
        return self.part.get_or_add_image_part(io.BytesIO(data))[1]

    def svg(self, svg: bytes, png_width: int) -> tuple:
        """An SVG picture with a PNG fallback: the ids of the PNG and the SVG."""
        png_rid = self.image(svg_to_png(svg, png_width))
        svg_part = Part(self.package.next_image_partname("svg"), "image/svg+xml", self.package, svg)
        return png_rid, self.part.relate_to(svg_part, RT.IMAGE)

    def icon(self, name: str, colour: str) -> tuple:
        return self.svg(lucide_svg(name, colour), 160)

    def chart(self, chart_type, data) -> tuple:
        """A native chart: python-pptx writes the chart part and the workbook behind it, for Edit Data."""
        rid = self.part.add_chart_part(chart_type, data)
        return rid, self.part.related_part(rid).chart


# ---------------------------------------------------------------------------------------------------------------------
# Theme, master and layouts

def theme_xml() -> str:
    colours = [("dk1", BLACK), ("lt1", WHITE), ("dk2", NAVY), ("lt2", OFFWHITE), ("accent1", NAVY),
               ("accent2", GREEN), ("accent3", GREEN50), ("accent4", CHARCOAL), ("accent5", LIME),
               ("accent6", GREY), ("hlink", GREEN), ("folHlink", CHARCOAL)]
    scheme = "".join(f'<a:{k}><a:srgbClr val="{v}"/></a:{k}>' for k, v in colours)
    fill = '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>'
    lines = "".join(f'<a:ln w="{w}" cap="flat" cmpd="sng" algn="ctr">{fill}<a:prstDash val="solid"/>'
                    '<a:miter lim="800000"/></a:ln>' for w in (9525, 12700, 19050))
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="BwB">'
            f'<a:themeElements><a:clrScheme name="BwB">{scheme}</a:clrScheme>'
            f'<a:fontScheme name="BwB"><a:majorFont><a:latin typeface="{HEAD}"/><a:ea typeface=""/><a:cs typeface=""/>'
            f'</a:majorFont><a:minorFont><a:latin typeface="{BODY}"/><a:ea typeface=""/><a:cs typeface=""/>'
            '</a:minorFont></a:fontScheme>'
            f'<a:fmtScheme name="BwB"><a:fillStyleLst>{fill * 3}</a:fillStyleLst><a:lnStyleLst>{lines}</a:lnStyleLst>'
            '<a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/>'
            '</a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst>'
            f'<a:bgFillStyleLst>{fill * 3}</a:bgFillStyleLst></a:fmtScheme></a:themeElements>'
            # new shapes: navy with white Aptos text; new lines: navy; new text boxes: Aptos 14 black, no margins
            '<a:objectDefaults><a:spDef><a:spPr><a:solidFill><a:schemeClr val="accent1"/></a:solidFill>'
            '<a:ln><a:noFill/></a:ln></a:spPr><a:bodyPr rtlCol="0" anchor="ctr" lIns="72000" tIns="72000" '
            'rIns="72000" bIns="72000"/><a:lstStyle><a:defPPr><a:defRPr lang="en-GB" sz="1400"><a:solidFill>'
            '<a:schemeClr val="lt1"/></a:solidFill></a:defRPr></a:defPPr></a:lstStyle></a:spDef>'
            '<a:lnDef><a:spPr><a:ln w="12700"><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:ln>'
            '</a:spPr><a:bodyPr/><a:lstStyle/></a:lnDef>'
            '<a:txDef><a:spPr><a:noFill/></a:spPr><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" '
            'rtlCol="0"><a:spAutoFit/></a:bodyPr><a:lstStyle><a:defPPr><a:defRPr lang="en-GB" sz="1400">'
            '<a:solidFill><a:schemeClr val="tx1"/></a:solidFill></a:defRPr></a:defPPr></a:lstStyle></a:txDef>'
            '</a:objectDefaults><a:extraClrSchemeLst/></a:theme>')


def table_styles_xml() -> str:
    """A BwB table style, the default for new tables: navy header row, light grey rules, total row in off-white."""
    grey = f'<a:ln w="9525" cmpd="sng">{solid(GREY)}</a:ln>'
    none = "<a:ln><a:noFill/></a:ln>"

    def txt(colour: str, bold: bool) -> str:
        b = ' b="on"' if bold else ""
        return (f'<a:tcTxStyle{b}><a:fontRef idx="minor"><a:prstClr val="black"/></a:fontRef>'
                f'<a:srgbClr val="{colour}"/></a:tcTxStyle>')

    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<a:tblStyleLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" def="{BWB_TABLE_STYLE}">'
            f'<a:tblStyle styleId="{BWB_TABLE_STYLE}" styleName="BwB table"><a:wholeTbl>{txt(BLACK, False)}'
            f'<a:tcStyle><a:tcBdr><a:left>{none}</a:left><a:right>{none}</a:right><a:top>{grey}</a:top>'
            f'<a:bottom>{grey}</a:bottom><a:insideH>{grey}</a:insideH><a:insideV>{none}</a:insideV></a:tcBdr>'
            '<a:fill><a:noFill/></a:fill></a:tcStyle></a:wholeTbl>'
            f'<a:lastRow>{txt(BLACK, True)}<a:tcStyle><a:tcBdr><a:top><a:ln w="12700" cmpd="sng">{solid(NAVY)}</a:ln>'
            f'</a:top></a:tcBdr><a:fill>{solid(OFFWHITE)}</a:fill></a:tcStyle></a:lastRow>'
            f'<a:firstRow>{txt(WHITE, True)}<a:tcStyle><a:tcBdr/><a:fill>{solid(NAVY)}</a:fill></a:tcStyle>'
            '</a:firstRow></a:tblStyle></a:tblStyleLst>')


def master_xml(logo_rid: str, logo_aspect: float) -> str:
    cv = Canvas()
    cv.placeholder("Title", 'type="title"', PH("Action title: say the so-what in one sentence"), G.title,
                   body_pr("b"))
    cv.placeholder("Text", 'type="body" idx="1"',
                   PH("Text in Aptos 14 pt: lead with the point") + PH("Second level: the supporting detail", 1)
                   + PH("Third level", 2), G.content, body_pr("t"))
    fx, fy, fw, fh = G.footer
    cv.slide_number('type="sldNum" sz="quarter" idx="4"', box=(fx + fw - 80, fy, 80, fh), bpr=body_pr("ctr"),
                    lst=lvl_ppr(1, 9, CHARCOAL, align="r"))
    logo_h = 16
    cv.picture("BwB logo", logo_rid, (fx, fy + (fh - logo_h) / 2, logo_h * logo_aspect, logo_h),
               "Bankers without Boundaries logo")
    cv.strip()
    layouts = "".join(f'<p:sldLayoutId id="{2147483649 + i}" r:id="rId{i + 1}"/>' for i in range(len(LAYOUT_NAMES)))
    title_style = lvl_ppr(1, 24, NAVY, "+mj-lt", bold=True, line=90)
    body_levels = [lvl_ppr(1, 14, BLACK, after=8, bullet="square"), lvl_ppr(2, 12, BLACK, after=6, bullet="circle")]
    body_levels += [lvl_ppr(n, 12, BLACK, after=4, bullet="dash") for n in range(3, 10)]
    other = "".join(lvl_ppr(n, 14, BLACK) for n in range(1, 10))
    guides = [("vert", G.margin_left), ("vert", G.x(6) - G.gutter), ("vert", G.x(6)), ("vert", RIGHT),
              ("horz", G.title[1]), ("horz", G.title[1] + G.title[3]), ("horz", CT), ("horz", CB),
              ("horz", G.source[1]), ("horz", G.footer[1])]
    guide_xml = "".join(f'<p15:guide id="{i + 1}" orient="{o}" pos="{int(round(v * 8))}" userDrawn="1"><p15:clr>'
                        '<a:srgbClr val="F26B43"/></p15:clr></p15:guide>' for i, (o, v) in enumerate(guides))
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sldMaster {NS}>'
            f'<p:cSld><p:bg><p:bgPr>{solid(WHITE)}<a:effectLst/></p:bgPr></p:bg>{sp_tree(cv.xml())}'
            f'{creation_ext("master")}</p:cSld>'
            '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" '
            'accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
            f'<p:sldLayoutIdLst>{layouts}</p:sldLayoutIdLst><p:hf hdr="0" ftr="0" dt="0"/>'
            f'<p:txStyles><p:titleStyle>{title_style}</p:titleStyle><p:bodyStyle>{"".join(body_levels)}</p:bodyStyle>'
            f'<p:otherStyle><a:defPPr><a:defRPr lang="en-GB"/></a:defPPr>{other}</p:otherStyle></p:txStyles>'
            f'<p:extLst><p:ext uri="{GUIDE_EXT}"><p15:sldGuideLst xmlns:p15="{P15}">{guide_xml}</p15:sldGuideLst>'
            "</p:ext></p:extLst></p:sldMaster>")


def layout_xml(name: str, white_logo_rid: str, white_logo_aspect: float) -> str:
    """One of the seven layouts. The navy layouts hide the master's footer and draw their own."""
    cv = Canvas()
    navy = name in NAVY_LAYOUTS
    slide_number = 'type="sldNum" sz="quarter" idx="12"'
    wide = G.w(10)

    def logo():
        cv.picture("BwB logo", white_logo_rid, (G.margin_left, G.margin_top, 48 * white_logo_aspect, 48),
                   "Bankers without Boundaries logo")

    if name == "BwB Title":
        logo()
        cv.placeholder("Title", 'type="ctrTitle" hasCustomPrompt="1"', PH("Presentation Title in Title Case"),
                       (G.margin_left, 168, wide, 168), body_pr("b"),
                       lvl_ppr(1, 40, WHITE, "+mj-lt", bold=True, line=90))
        cv.shape("Lime rule", (G.margin_left, 352, G.w(2), 4), fill=LIME)
        cv.placeholder("Subtitle", 'type="subTitle" idx="1" hasCustomPrompt="1"',
                       PH("Subtitle: the client, the project and the purpose of the meeting"),
                       (G.margin_left, 372, wide, 56), body_pr("t"), lvl_ppr(1, 18, WHITE))
        cv.placeholder("Date", 'type="body" sz="quarter" idx="13" hasCustomPrompt="1"', PH("Date: 5 October 2026"),
                       (G.margin_left, 440, G.w(6), 24), body_pr("t"), lvl_ppr(1, 14, WHITE))
        cv.strip()
    elif name == "BwB Section divider":
        cv.placeholder("Section number", 'type="body" sz="quarter" idx="13" hasCustomPrompt="1"', PH("01"),
                       (G.margin_left, 96, G.w(6), 160), body_pr("b"),
                       lvl_ppr(1, 120, LIME, "+mj-lt", bold=True, line=80))
        cv.placeholder("Title", 'type="title" hasCustomPrompt="1"', PH("Section title in sentence case"),
                       (G.margin_left, 280, wide, 96), body_pr("t"),
                       lvl_ppr(1, 36, WHITE, "+mj-lt", bold=True, line=90))
        cv.placeholder("Description", 'type="body" sz="quarter" idx="14" hasCustomPrompt="1"',
                       PH("One sentence on what this section shows"), (G.margin_left, 392, G.w(8), 48),
                       body_pr("t"), lvl_ppr(1, 16, WHITE))
        cv.strip()
    elif name == "BwB Closing":
        logo()
        cv.placeholder("Title", 'type="title" hasCustomPrompt="1"', PH("Thank you"),
                       (G.margin_left, 168, wide, 96), body_pr("b"),
                       lvl_ppr(1, 40, WHITE, "+mj-lt", bold=True, line=90))
        cv.shape("Lime rule", (G.margin_left, 280, G.w(2), 4), fill=LIME)
        for i in (1, 2):
            cv.placeholder(f"Contact {i}", f'type="body" idx="{i}" hasCustomPrompt="1"',
                           PH("**Name Surname**") + PH("Role") + PH("name.surname@bwb.earth"),
                           (G.x(4 * (i - 1)), 304, G.w(4), 120), body_pr("t"), lvl_ppr(1, 14, WHITE, after=2))
        cv.text("Website", (G.margin_left, 476, G.w(4), 20), P("bwb.earth", 14, WHITE, bold=True), anchor="b")
        cv.strip()
    else:
        if name != "BwB Blank":
            cv.placeholder("Title", 'type="title" hasCustomPrompt="1"',
                           PH("Action title: say the so-what in one sentence, on no more than two lines"))
        prompt = PH("Text in Aptos 14 pt: lead with the point, then the evidence") + PH("Second level", 1)
        if name == "BwB Title and content":
            cv.placeholder("Content", 'idx="1" hasCustomPrompt="1"', prompt)
        elif name == "BwB Two content":
            half = G.w(6)
            cv.placeholder("Left content", 'sz="half" idx="1" hasCustomPrompt="1"', prompt, (G.x(0), CT, half, CH))
            cv.placeholder("Right content", 'sz="half" idx="2" hasCustomPrompt="1"', prompt, (G.x(6), CT, half, CH))
        cv.slide_number(slide_number)
    kind = {"BwB Title": "title", "BwB Section divider": "secHead", "BwB Title and content": "obj",
            "BwB Two content": "twoObj", "BwB Title only": "titleOnly", "BwB Blank": "blank"}.get(name)
    attrs = (f' type="{kind}"' if kind else "") + (' showMasterSp="0"' if navy else "")
    bg = f"<p:bg><p:bgPr>{solid(NAVY)}<a:effectLst/></p:bgPr></p:bg>" if navy else ""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sldLayout {NS}{attrs} preserve="1" '
            f'userDrawn="1"><p:cSld name="{name}">{bg}{sp_tree(cv.xml())}{creation_ext("layout/" + name)}</p:cSld>'
            "<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>")


def core_xml(title: str) -> str:
    """Document properties. The starter deck has no title, author or dates, so a deck made from it starts clean."""
    props = ""
    if title:
        props = (f"<dc:title>{esc(title)}</dc:title><dc:creator>Bankers without Boundaries</dc:creator>"
                 "<cp:lastModifiedBy>Bankers without Boundaries</cp:lastModifiedBy>"
                 f'<dcterms:created xsi:type="dcterms:W3CDTF">{FIXED_ISO}</dcterms:created>'
                 f'<dcterms:modified xsi:type="dcterms:W3CDTF">{FIXED_ISO}</dcterms:modified>')
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f"{props}<cp:revision>1</cp:revision><dc:language>en-GB</dc:language></cp:coreProperties>")


def base_package(slide_count: int, title: str) -> bytes:
    """A pptx with the theme, the master and the seven layouts and no slides, for python-pptx to open."""
    footer_logo, footer_aspect = trimmed_png(BRAND / "header-logo.png")
    white_logo, white_aspect = trimmed_png(BRAND / "logo-white.png")
    rel = ('<Relationship Id="{}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{}" '
           'Target="{}"/>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{}</Relationships>')
    ct = "application/vnd.openxmlformats-officedocument.presentationml."
    overrides = [("/ppt/presentation.xml", ct + "presentation.main+xml"),
                 ("/ppt/slideMasters/slideMaster1.xml", ct + "slideMaster+xml"),
                 ("/ppt/theme/theme1.xml", "application/vnd.openxmlformats-officedocument.theme+xml"),
                 ("/ppt/presProps.xml", ct + "presProps+xml"), ("/ppt/viewProps.xml", ct + "viewProps+xml"),
                 ("/ppt/tableStyles.xml", ct + "tableStyles+xml"),
                 ("/docProps/core.xml", "application/vnd.openxmlformats-package.core-properties+xml"),
                 ("/docProps/app.xml", "application/vnd.openxmlformats-officedocument.extended-properties+xml")]
    overrides += [(f"/ppt/slideLayouts/slideLayout{i + 1}.xml", ct + "slideLayout+xml")
                  for i in range(len(LAYOUT_NAMES))]
    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/>'
            + "".join(f'<Override PartName="{p}" ContentType="{c}"/>' for p, c in overrides) + "</Types>"),
        "_rels/.rels": rels.format(
            rel.format("rId1", "officeDocument", "ppt/presentation.xml")
            + '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/'
              'core-properties" Target="docProps/core.xml"/>'
            + rel.format("rId3", "extended-properties", "docProps/app.xml")),
        "docProps/core.xml": core_xml(title),
        "docProps/app.xml": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Properties '
            'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
            'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            "<Application>Microsoft Office PowerPoint</Application><PresentationFormat>Widescreen</PresentationFormat>"
            f"<Slides>{slide_count}</Slides><Company>Bankers without Boundaries</Company>"
            "<AppVersion>16.0000</AppVersion></Properties>"),
        "ppt/_rels/presentation.xml.rels": rels.format(
            rel.format("rId1", "slideMaster", "slideMasters/slideMaster1.xml")
            + rel.format("rId2", "theme", "theme/theme1.xml")
            + rel.format("rId3", "presProps", "presProps.xml") + rel.format("rId4", "viewProps", "viewProps.xml")
            + rel.format("rId5", "tableStyles", "tableStyles.xml")),
        "ppt/presentation.xml": (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:presentation {NS}>'
            '<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst><p:sldIdLst/>'
            f'<p:sldSz cx="{emu(G.width)}" cy="{emu(G.height)}"/><p:notesSz cx="6858000" cy="9144000"/>'
            f'<p:defaultTextStyle><a:defPPr><a:defRPr lang="en-GB"/></a:defPPr>'
            + "".join(lvl_ppr(n, 14, BLACK) for n in range(1, 10)) + "</p:defaultTextStyle></p:presentation>"),
        "ppt/presProps.xml": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:presentationPr {NS}/>',
        "ppt/viewProps.xml": (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:viewPr {NS}><p:normalViewPr>'
            '<p:restoredLeft sz="15620"/><p:restoredTop sz="94660"/></p:normalViewPr><p:slideViewPr><p:cSldViewPr>'
            '<p:cViewPr varScale="1"><p:scale><a:sx n="100" d="100"/><a:sy n="100" d="100"/></p:scale>'
            '<p:origin x="0" y="0"/></p:cViewPr><p:guideLst/></p:cSldViewPr></p:slideViewPr><p:gridSpacing '
            f'cx="{emu(4)}" cy="{emu(4)}"/></p:viewPr>'),
        "ppt/tableStyles.xml": table_styles_xml(),
        "ppt/theme/theme1.xml": theme_xml(),
        "ppt/media/image1.png": footer_logo,
        "ppt/media/image2.png": white_logo,
        "ppt/slideMasters/slideMaster1.xml": master_xml(f"rId{len(LAYOUT_NAMES) + 2}", footer_aspect),
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": rels.format(
            "".join(rel.format(f"rId{i + 1}", "slideLayout", f"../slideLayouts/slideLayout{i + 1}.xml")
                    for i in range(len(LAYOUT_NAMES)))
            + rel.format(f"rId{len(LAYOUT_NAMES) + 1}", "theme", "../theme/theme1.xml")
            + rel.format(f"rId{len(LAYOUT_NAMES) + 2}", "image", "../media/image1.png")),
    }
    for i, name in enumerate(LAYOUT_NAMES):
        files[f"ppt/slideLayouts/slideLayout{i + 1}.xml"] = layout_xml(name, "rId2", white_aspect)
        logo = rel.format("rId2", "image", "../media/image2.png") if name in WHITE_LOGO_LAYOUTS else ""
        files[f"ppt/slideLayouts/_rels/slideLayout{i + 1}.xml.rels"] = rels.format(
            rel.format("rId1", "slideMaster", "../slideMasters/slideMaster1.xml") + logo)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------------------------------
# The library slides. Each builder draws one slide; its text is written as guidance for the author.

def read_disclaimer() -> tuple:
    """The standard disclaimer and the FCA note, from src/docx/sections.ts, so slides and documents match."""
    ts = (ROOT / "src" / "docx" / "sections.ts").read_text(encoding="utf-8")
    block = re.search(r"export const DEFAULT_DISCLAIMER: readonly string\[\] = \[(.*?)\n\]", ts, re.S)
    note = re.search(r"export const DISCLAIMER_NOTE = '((?:[^'\\]|\\.)*)'", ts)
    if not block or not note:
        sys.exit("Could not read DEFAULT_DISCLAIMER and DISCLAIMER_NOTE from src/docx/sections.ts")

    def unescape(text: str) -> str:
        text = re.sub(r"\\u\{([0-9A-Fa-f]+)\}", lambda m: chr(int(m.group(1), 16)), text)
        return text.replace("\\'", "'")

    paras = [unescape(t) for t in re.findall(r"'((?:[^'\\]|\\.)*)'", block.group(1))]
    return paras, unescape(note.group(1))


def title(cv: Canvas, text: str):
    cv.placeholder("Title", 'type="title"', PH(text))


def source(cv: Canvas, text: str):
    cv.text("Source", G.source, P(text, 8, CHARCOAL), anchor="b")


def callout(cv: Canvas, box: tuple, heading: str, text: str, name: str = "Callout"):
    """An off-white box with a green bar, as the Callout element in the BwB tab."""
    x, y, w, h = box
    cv.shape(name, box, fill=OFFWHITE, text=P(heading, 14, NAVY, bold=True, after=4) + P(text, 12), anchor="t",
             inset=(20, 12, 16, 12))
    cv.shape(f"{name} bar", (x, y, 4, h), fill=GREEN)


def coloured(parts: list, size: float, align: str = "l") -> str:
    """One paragraph of bold Aptos runs, each in its own colour, such as a lime step number on navy."""
    out = "".join(f"<a:r>{rpr(size, colour, BODY, True)}<a:t>{esc(text)}</a:t></a:r>" for text, colour in parts)
    return (f'<a:p><a:pPr marL="0" indent="0" algn="{align}"><a:buNone/></a:pPr>{out}'
            f'{rpr(size, parts[-1][1], BODY, tag="a:endParaRPr")}</a:p>')


def figure(value: str, size: float, colour: str) -> str:
    """A key figure in Aptos bold. A leading currency code is set smaller, as in the BwB documents."""
    m = re.match(r"([A-Z]{3}) (.+)", value)
    out = f'<a:r>{rpr(round(size * 0.42), colour, BODY, True)}<a:t>{m.group(1)} </a:t></a:r>' if m else ""
    out += f"<a:r>{rpr(size, colour, BODY, True)}<a:t>{esc(m.group(2) if m else value)}</a:t></a:r>"
    return (f'<a:p><a:pPr marL="0" indent="0" algn="l"><a:lnSpc><a:spcPct val="90000"/></a:lnSpc><a:buNone/></a:pPr>'
            f'{out}{rpr(size, colour, BODY, tag="a:endParaRPr")}</a:p>')


def s_title(cv: Canvas, a: Assets):
    cv.placeholder("Title", 'type="ctrTitle"', PH("Presentation Title in Title Case, on One or Two Lines"))
    cv.placeholder("Subtitle", 'type="subTitle" idx="1"',
                   PH("Subtitle: the client, the project and the purpose of the meeting"))
    cv.placeholder("Date", 'type="body" sz="quarter" idx="13"', PH("Date in the form 5 October 2026"))


AGENDA = [("Context and objectives", "Why we are here and the decision this meeting supports"),
          ("Findings", "What the diagnostic shows, with the evidence behind each point"),
          ("Options", "The options we assessed and how they compare"),
          ("Recommendations", "What we recommend, why, and what it will take"),
          ("Next steps", "Who does what by when, and the decisions we need today")]


def s_agenda(cv: Canvas, a: Assets):
    title(cv, "Agenda")
    row = CH / len(AGENDA)
    cv.line("Agenda rule top", G.margin_left, CT, RIGHT, CT, NAVY, 1)
    for i, (head, text) in enumerate(AGENDA):
        y = CT + i * row
        cv.text(f"Agenda number {i + 1}", (G.x(0), y, G.w(1), row), P(f"{i + 1:02d}", 24, GREEN, BODY, bold=True),
                anchor="ctr")
        cv.text(f"Agenda item {i + 1}", (G.x(1), y, G.w(5), row), P(head, 18, NAVY, HEAD, bold=True), anchor="ctr")
        cv.text(f"Agenda description {i + 1}", (G.x(6), y, G.w(6), row), P(text, 14, CHARCOAL), anchor="ctr")
        cv.line(f"Agenda rule {i + 1}", G.margin_left, y + row, RIGHT, y + row, GREY)


def s_divider(cv: Canvas, a: Assets):
    cv.placeholder("Section number", 'type="body" sz="quarter" idx="13"', PH("01"))
    cv.placeholder("Title", 'type="title"', PH("Section title in sentence case, as it appears on the agenda"))
    cv.placeholder("Description", 'type="body" sz="quarter" idx="14"',
                   PH("One sentence on what this section shows and why it matters to the reader"))


SUMMARY = [("State the answer to the client\u2019s question in one sentence",
            ["Put the conclusion first: a reader who stops here should still know what we recommend",
             "Use the client\u2019s own words for the question"]),
           ("Give the finding that matters most, with its evidence",
            ["Back it with one or two facts: a figure, a comparison or a trend",
             "Source each fact on the slide that shows it in full"]),
           ("Close with what the client should do next",
            ["Name the decision, the owner and the timing",
             "Keep each message to two lines, so the page reads in a minute"])]


def s_summary(cv: Canvas, a: Assets):
    title(cv, "Executive summary")
    gap = 16
    row = (CH - 2 * gap) / 3
    for i, (head, points) in enumerate(SUMMARY):
        y = CT + i * (row + gap)
        cv.shape(f"Message marker {i + 1}", (G.x(0), y, 32, 32), fill=GREEN, anchor="ctr",
                 text=P(str(i + 1), 16, WHITE, BODY, bold=True, align="ctr"))
        cv.text(f"Message {i + 1}", (G.x(1), y, G.w(5), row), P(head, 18, NAVY, HEAD, bold=True, line=95))
        cv.text(f"Message detail {i + 1}", (G.x(6), y + 2, G.w(6), row),
                "".join(P(t, 14, bullet="square", after=6) for t in points))
        if i < 2:
            cv.line(f"Message rule {i + 1}", G.margin_left, y + row + gap / 2, RIGHT, y + row + gap / 2, GREY)


def s_bullets(cv: Canvas, a: Assets):
    title(cv, "Action title: say the so-what in one sentence, on no more than two lines")
    items = [(0, "Lead with the point, then give the evidence for it"),
             (1, "Use the second level for supporting detail, at most three under a point"),
             (0, "Keep to five or six bullets of one or two lines each"),
             (0, "Write full sentences in plain English, with UK spelling and no filler"),
             (1, "Spell out an abbreviation the first time it appears"),
             (0, "Use figures where they make the point: GBP 5m, 38 per cent, 2025 to 2030"),
             (1, "Give the source of every figure in the source line at the foot of the slide"),
             (0, "Make the bullets parallel: start each with the same kind of word"),
             (0, "Bold only the **few words** the reader must not miss")]
    cv.placeholder("Content", 'idx="1"', "".join(PH(t, lvl) for lvl, t in items), (G.x(0), CT, G.w(8), CH))
    callout(cv, (G.x(8), CT, G.w(4), 92), "Key takeaway",
            "Use a callout for the one message the slide exists to make. Keep it to three lines.")
    source(cv, "Source: organisation, title of the publication (year). Notes go on the same line, in Aptos 8 pt "
               "charcoal.")


def s_two_columns(cv: Canvas, a: Assets):
    title(cv, "Action title: compare two options, periods or groups side by side")
    cols = [("Option A: name it in a few words",
             ["Say what the option involves, in one line", "Give its main benefit, with a figure if there is one",
              "Give its main risk or cost", "Say what it would take to deliver"]),
            ("Option B: name it in a few words",
             ["Match the left column point for point, so the reader can compare across",
              "Use the same units and periods in both columns", "Keep the points parallel and about the same length",
              "Leave the verdict to the box below"])]
    box = 64
    for i, (head, items) in enumerate(cols):
        x, w = G.x(6 * i), G.w(6)
        cv.text(f"Column heading {i + 1}", (x, CT, w, 28), P(head, 18, NAVY, HEAD, bold=True), anchor="b")
        cv.line(f"Column rule {i + 1}", x, CT + 38, x + w, CT + 38, GREEN, 1.5)
        cv.placeholder("Left content" if i == 0 else "Right content", f'sz="half" idx="{i + 1}"',
                       "".join(PH(t) for t in items), (x, CT + 56, w, CH - 56 - box - 16))
    callout(cv, (G.x(0), CB - box, G.w(12), box), "Recommendation",
            "Say which option we recommend and the one reason that decides it.")


FIGURES = [("USD 250m", "Headline figure, with its unit and year",
            "Lead with the figure that carries the message. Only this one is lime."),
           ("38%", "Second figure: say what it measures",
            "Round figures to what the reader needs: 38%, not 37.84%."),
           ("12", "Countries, projects or partners",
            "Use the % sign in figures and \u2018per cent\u2019 in sentences."),
           ("2030", "A date or a target year",
            "Write currency as a code and a figure: GBP 5m, USD 250m, EUR 1.2bn.")]


def s_key_figures(cv: Canvas, a: Assets):
    title(cv, "Action title: what the numbers show, in one sentence")
    cv.text("Introduction", (G.x(0), CT, G.w(8), 40),
            P("One or two lines that set up the figures: the scope, the period and where the data comes from.", 14))
    top, band = CT + 56, 176
    cv.shape("Navy band", (0, top, G.width, band), fill=NAVY)
    for i, (value, label, note) in enumerate(FIGURES):
        x, w = G.x(3 * i), G.w(3)
        if i:
            cv.line(f"Key figure divider {i}", x - G.gutter / 2, top + 28, x - G.gutter / 2, top + band - 28, WHITE,
                    0.5)
        cv.text(f"Key figure {i + 1}", (x, top + 24, w, 64), figure(value, 44, LIME if i == 0 else WHITE), anchor="b")
        cv.text(f"Key figure label {i + 1}", (x, top + 100, w, 48), P(label, 14, WHITE))
        cv.text(f"Key figure note {i + 1}", (x, top + band + 20, w, CB - top - band - 20), P(note, 12))
    source(cv, "Source: organisation, title of the publication (year). Every figure needs a source.")


CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def style_chart(chart):
    """BwB chart style: green and green 50 series, Aptos labels in charcoal, no title, no gridlines, no frame."""
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.font.name = BODY
    chart.font.size = Pt(10)
    chart.font.color.rgb = RGBColor.from_string(CHARCOAL)
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -8
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.number_format = "#,##0"
    labels.number_format_is_linked = False
    labels.font.size = Pt(10)
    labels.font.color.rgb = RGBColor.from_string(CHARCOAL)
    for series, colour in zip(plot.series, (GREEN, GREEN50)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(colour)
        series.format.line.fill.background()
    value = chart.value_axis
    value.visible = False
    value.has_major_gridlines = False
    value.major_tick_mark = XL_TICK_MARK.NONE
    category = chart.category_axis
    category.major_tick_mark = XL_TICK_MARK.NONE
    category.format.line.color.rgb = RGBColor.from_string(GREY20)
    category.tick_labels.font.size = Pt(10)
    category.tick_labels.font.color.rgb = RGBColor.from_string(CHARCOAL)
    space = chart._chartSpace

    def tag(name: str) -> str:
        return f"{{{CHART_NS}}}{name}"

    for axis in space.iter(tag("axId"), tag("crossAx")):
        axis.set("val", axis.get("val", "").lstrip("-"))  # python-pptx writes negative ids; the schema wants unsigned
    for hidden in space.iter(tag("delete")):
        hidden.set("val", hidden.get("val", "1"))  # say "hidden" outright rather than rely on the schema default
    lang = space.find(tag("lang"))
    if lang is not None:
        lang.set("val", "en-GB")
    if space.find(tag("roundedCorners")) is None:
        anchor = lang if lang is not None else space.find(tag("date1904"))
        corners = parse_xml(f'<c:roundedCorners xmlns:c="{CHART_NS}" val="0"/>')
        if anchor is not None:
            anchor.addnext(corners)
        else:
            space.insert(0, corners)
    if space.find(tag("spPr")) is None:
        space.find(tag("chart")).addnext(parse_xml(
            f'<c:spPr xmlns:c="{CHART_NS}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            "<a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>"))


def s_chart(cv: Canvas, a: Assets):
    title(cv, "Action title: the conclusion the chart supports, not a description of it")
    w8 = G.w(8)
    cv.text("Chart heading", (G.x(0), CT, w8, 20),
            P("Chart heading: what is measured, the unit and the period (USD m, 2021 to 2025)", 12, NAVY, bold=True))
    data = CategoryChartData(number_format="#,##0")
    data.categories = ["2021", "2022", "2023", "2024", "2025"]
    data.add_series("Committed", (120, 150, 185, 210, 260))
    data.add_series("Disbursed", (80, 105, 140, 170, 205))
    rid, chart = a.chart(XL_CHART_TYPE.COLUMN_CLUSTERED, data)
    style_chart(chart)
    cv.chart("Chart", rid, (G.x(0), CT + 28, w8, CH - 28),
             "Column chart of two series over five years. Right-click it and choose Edit Data to use your own figures")
    x, w = G.x(8), G.w(4)
    points = ["Name the pattern: the trend, the gap or the outlier", "Give the number that proves it, read from the chart",
              "Say what it means for the decision in hand",
              "Right-click the chart and choose Edit Data to put in your own figures"]
    cv.shape("Commentary", (x, CT, w, CH), fill=OFFWHITE, anchor="t", inset=(20, 20, 20, 20),
             text=P("What the chart shows", 16, NAVY, HEAD, bold=True, after=12)
             + "".join(P(t, 12, bullet="square", after=8) for t in points))
    source(cv, "Source: organisation, dataset (year). Charts use green and green 50 on white, with no title inside "
               "the chart.")


TABLE = [("Green bonds", "1,250", "1,480", "18", "Figures right-aligned, with a thousands separator"),
         ("Sustainability-linked loans", "860", "910", "6", "Units in the column heading, not in each cell"),
         ("Blended finance", "420", "510", "21", "The % sign in tables, \u2018per cent\u2019 in sentences"),
         ("Guarantees", "180", "240", "33", "Light grey rules between rows, no vertical lines"),
         ("Results-based finance", "95", "120", "26", "Order the rows so that the order means something"),
         ("Total", "2,805", "3,260", "16", "Totals in bold, on the off-white row")]


def s_table(cv: Canvas, a: Assets):
    title(cv, "Action title: the comparison the table makes, in one sentence")
    num = 112
    widths = [G.w(3), num, num, num, G.width - G.margin_left - G.margin_right - G.w(3) - 3 * num]
    head = ["Instrument", "2024 (USD m)", "2025 (USD m)", "Change (%)", "How the table is set"]
    rows = [(44, [tc(P(t, 12, WHITE, bold=True, align="r" if 0 < j < 4 else "l"), NAVY, margins=(8, 4, 8, 4))
                  for j, t in enumerate(head)])]
    for i, cells in enumerate(TABLE):
        total = i == len(TABLE) - 1
        lines = {"T": (NAVY, 1), "B": (GREY, 0.75)} if total else {"B": (GREY, 0.75)}
        rows.append((44, [tc(P(t, 12, CHARCOAL if j == 4 else BLACK, bold=total and j < 4,
                               align="r" if 0 < j < 4 else "l"),
                             OFFWHITE if total else None, lines, margins=(8, 4, 8, 4)) for j, t in enumerate(cells)]))
    cv.table("Table", G.margin_left, CT, widths, rows)
    source(cv, "Source: organisation, dataset (year). Illustrative figures.")


STEPS = [("Diagnose", "Weeks 1 to 4", ["Frame the question and the decision it informs",
                                       "Collect the data and interview stakeholders",
                                       "Test early findings with the client team"], "Diagnostic note"),
         ("Design", "Weeks 5 to 8", ["Develop the options and test them with the client",
                                     "Agree the criteria for choosing between them",
                                     "Shortlist two or three options to detail"], "Options paper"),
         ("Structure", "Weeks 9 to 12", ["Detail the preferred option: instruments, terms and risks",
                                         "Model the costs and the benefits", "Draft the term sheet for investors"],
          "Structuring report"),
         ("Mobilise", "Weeks 13 to 16", ["Engage investors and partners", "Plan the launch and the first 100 days",
                                         "Set up the reporting on results"], "Implementation plan")]


def s_process(cv: Canvas, a: Assets):
    title(cv, "Action title: how the work moves from question to answer in four steps")
    depth, h, pad = 16, 52, 14
    for i, (step, when, items, output) in enumerate(STEPS):
        x, w = G.x(3 * i), G.w(3)
        reach = G.gutter - 4 if i < len(STEPS) - 1 else 0
        cv.shape(f"Step {i + 1}", (x, CT, w + reach, h), geom="homePlate" if i == 0 else "chevron",
                 adj={"adj": int(depth / h * 100000)}, fill=NAVY, anchor="ctr", inset=(depth, 0, depth, 0),
                 text=coloured([(f"{i + 1} ", LIME), (step, WHITE)], 16, "ctr"))
        top = CT + h + 12
        cv.shape(f"Step panel {i + 1}", (x, top, w, CB - top), fill=OFFWHITE)
        cv.text(f"Step timing {i + 1}", (x + pad, top + 14, w - 2 * pad, 18), P(when, 12, CHARCOAL, bold=True))
        cv.text(f"Step text {i + 1}", (x + pad, top + 42, w - 2 * pad, 170),
                "".join(P(t, 12, bullet="square", after=6) for t in items))
        cv.line(f"Step rule {i + 1}", x + pad, CB - 50, x + w - pad, CB - 50, GREY20)
        cv.text(f"Step output {i + 1}", (x + pad, CB - 40, w - 2 * pad, 30), P(f"**Output:** {output}", 12, NAVY))


QUADRANTS = [("Plan carefully", "High impact, but it needs preparation"),
             ("Prioritise", "High impact and ready to start"),
             ("Deprioritise", "Little gain for the effort"),
             ("Quick wins", "Easy to do, with modest gains")]
OPTIONS = [("A", "Option A: name it and say what it involves", (0.86, 0.33)),
           ("B", "Option B: one line each, in the same form", (0.66, 0.40)),
           ("C", "Option C: high impact, but it needs work first", (0.30, 0.36)),
           ("D", "Option D: easy, with a modest gain", (0.82, 0.82)),
           ("E", "Option E: say why it is not worth pursuing", (0.20, 0.84))]


def s_matrix(cv: Canvas, a: Assets):
    title(cv, "Action title: which options to pursue first, by impact and feasibility")
    x0, x1 = G.x(0) + 40, G.x(8) - G.gutter
    y0, y1 = CT, CB - 40
    gap = 8
    qw, qh = (x1 - x0 - gap) / 2, (y1 - y0 - gap) / 2
    for i, (head, text) in enumerate(QUADRANTS):
        qx, qy = x0 + (i % 2) * (qw + gap), y0 + (i // 2) * (qh + gap)
        hot = i == 1
        cv.shape(f"Quadrant {i + 1}", (qx, qy, qw, qh), fill=NAVY if hot else OFFWHITE, anchor="t",
                 inset=(16, 14, 16, 14), text=P(head, 16, WHITE if hot else NAVY, HEAD, bold=True, after=2)
                 + P(text, 12, WHITE if hot else BLACK))
    cv.line("Impact axis", G.x(0) + 28, y1, G.x(0) + 28, y0, CHARCOAL, 1, arrow=True)
    cv.text("Impact label", (G.x(0), y0, 20, y1 - y0), P("Impact", 12, CHARCOAL, bold=True, align="ctr"),
            anchor="ctr", vert="vert270")
    cv.line("Feasibility axis", x0, y1 + 12, x1, y1 + 12, CHARCOAL, 1, arrow=True)
    cv.text("Feasibility label", (x0, y1 + 18, x1 - x0, 20), P("Feasibility", 12, CHARCOAL, bold=True, align="ctr"))
    d = 24
    for letter, _, (fx, fy) in OPTIONS:
        cx, cy = x0 + fx * (x1 - x0), y0 + fy * (y1 - y0)
        cv.shape(f"Option {letter}", (cx - d / 2, cy - d / 2, d, d), geom="ellipse", fill=GREEN, line=WHITE,
                 line_w=1.5, anchor="ctr", text=P(letter, 12, WHITE, bold=True, align="ctr"))
    x, w = G.x(8), G.w(4)
    cv.text("Options heading", (x, CT, w, 24), P("Options assessed", 16, NAVY, HEAD, bold=True))
    for i, (letter, text, _) in enumerate(OPTIONS):
        y = CT + 48 + i * 52
        cv.shape(f"Option key {letter}", (x, y, d, d), geom="ellipse", fill=GREEN, anchor="ctr",
                 text=P(letter, 12, WHITE, bold=True, align="ctr"))
        cv.text(f"Option text {letter}", (x + 36, y + 3, w - 36, 44), P(text, 12))


PHASES = ["Phase 1: diagnostic", "Phase 2: design", "Phase 3: delivery"]
MILESTONES = [("January 2027", "Kick-off", "Inception meeting and data request", True),
              ("February 2027", "Inception report", "Approach, workplan and risks agreed", True),
              ("April 2027", "Diagnostic note", "Findings shared with the steering group", False),
              ("June 2027", "Options workshop", "Options tested with stakeholders", False),
              ("August 2027", "Draft report", "Recommendations shared for comment", False),
              ("October 2027", "Final report", "Final version and handover", False)]


def s_timeline(cv: Canvas, a: Assets):
    title(cv, "Action title: the milestones that matter between now and completion")
    for i, phase in enumerate(PHASES):
        cv.shape(f"Phase {i + 1}", (G.x(4 * i), CT, G.w(4), 28), fill=OFFWHITE, anchor="ctr", inset=(10, 0, 10, 0),
                 text=P(phase, 12, NAVY, bold=True))
    line_y = CT + 116
    cv.line("Timeline", G.margin_left, line_y, RIGHT, line_y, NAVY, 1.5)
    d = 14
    for i, (date, head, text, done) in enumerate(MILESTONES):
        x, w = G.x(2 * i), G.w(2)
        cv.text(f"Milestone date {i + 1}", (x, line_y - 40, w, 22), P(date, 12, CHARCOAL, bold=True), anchor="b")
        cv.shape(f"Milestone {i + 1}", (x, line_y - d / 2, d, d), geom="ellipse", fill=GREEN if done else WHITE,
                 line=GREEN, line_w=2)
        cv.text(f"Milestone title {i + 1}", (x, line_y + 22, w, 20), P(head, 14, NAVY, bold=True))
        cv.text(f"Milestone text {i + 1}", (x, line_y + 46, w, 60), P(text, 12))
    ly = line_y + 128
    cv.shape("Legend completed", (G.margin_left, ly, 10, 10), geom="ellipse", fill=GREEN, line=GREEN, line_w=1.5)
    cv.text("Legend completed text", (G.margin_left + 16, ly - 3, 80, 16), P("Completed", 10, CHARCOAL))
    cv.shape("Legend planned", (G.margin_left + 100, ly, 10, 10), geom="ellipse", fill=WHITE, line=GREEN, line_w=1.5)
    cv.text("Legend planned text", (G.margin_left + 116, ly - 3, 80, 16), P("Planned", 10, CHARCOAL))
    callout(cv, (G.x(0), CB - 64, G.w(12), 64), "Critical path",
            "Name the dependency that would delay the final milestone, who owns it and when it must be resolved.")


WORKPLAN = [("1 Inception", None, None),
            ("Kick-off meeting and data request", (1, 1), None),
            ("Inception report", (2, 3), 3),
            ("2 Diagnostic", None, None),
            ("Desk review and stakeholder interviews", (3, 7), None),
            ("Diagnostic note", (7, 8), 8),
            ("3 Options and recommendations", None, None),
            ("Options analysis and financial modelling", (8, 10), None),
            ("Validation workshop", (11, 11), 11),
            ("Final report", (11, 12), 12)]


def s_workplan(cv: Canvas, a: Assets):
    title(cv, "Action title: the workplan delivers the final report in week 12")
    weeks = 12
    first = G.x(4) - G.margin_left
    week_w = (G.width - G.margin_left - G.margin_right - first) / weeks
    white = (WHITE, 1)
    head1 = [tc(P("Workstream and activity", 11, WHITE, bold=True), NAVY, {"B": white}, rowspan=2, anchor="b",
                margins=(8, 4, 8, 6))]
    for m in range(3):
        head1.append(tc(P(f"Month {m + 1}", 11, WHITE, bold=True, align="ctr"), NAVY, {"L": white, "B": white},
                        span=4))
        head1 += [tc(P(""), NAVY, hmerge=True) for _ in range(3)]
    head2 = [tc(P(""), NAVY, vmerge=True)]
    head2 += [tc(P(str(n), 10, WHITE, align="ctr"), NAVY, {"L": white} if n in (1, 5, 9) else None,
                 margins=(0, 2, 0, 2)) for n in range(1, weeks + 1)]
    rows = [(24, head1), (22, head2)]
    for label, span, due in WORKPLAN:
        group = span is None
        cells = [tc(P(label, 11, NAVY if group else BLACK, bold=group), OFFWHITE if group else None,
                    {"B": (GREY, 0.75)}, margins=(8 if group else 22, 3, 8, 3))]
        for n in range(1, weeks + 1):
            fill = OFFWHITE if group else None
            if span and span[0] <= n <= span[1]:
                fill = NAVY if n == due else GREEN
            lines = {"B": (GREY, 0.75)}
            if n in (5, 9) and fill not in (GREEN, NAVY):
                lines["L"] = (GREY, 0.75)
            cells.append(tc(P(""), fill, lines, margins=(0, 3, 0, 3)))
        rows.append((28, cells))
    cv.table("Workplan", G.margin_left, CT, [first] + [week_w] * weeks, rows)
    ly = CT + 24 + 22 + 28 * len(WORKPLAN) + 14
    cv.shape("Legend activity", (G.margin_left, ly, 16, 10), fill=GREEN)
    cv.text("Legend activity text", (G.margin_left + 22, ly - 3, 80, 16), P("Activity", 10, CHARCOAL))
    cv.shape("Legend deliverable", (G.margin_left + 100, ly, 16, 10), fill=NAVY)
    cv.text("Legend deliverable text", (G.margin_left + 122, ly - 3, 120, 16), P("Deliverable due", 10, CHARCOAL))


ICONS = [("landmark", "Public finance", "One sentence on what the capability covers and who it serves."),
         ("leaf", "Nature and climate", "Choose each icon to name the topic: icons are never decoration."),
         ("handshake", "Partnerships", "Use Lucide icons in green on white, all at the same size."),
         ("trending-up", "Capital markets", "Keep titles to two or three words, in sentence case."),
         ("shield-check", "Risk and safeguards", "Keep the sentences parallel and about the same length."),
         ("graduation-cap", "Capacity building", "Six items at most; use three or four when each needs more text.")]


def s_icons(cv: Canvas, a: Assets):
    title(cv, "Action title: six capabilities that cover the client\u2019s needs end to end")
    gap = 16
    rh = (CH - gap) / 2
    for i, (icon, head, text) in enumerate(ICONS):
        x, w, y = G.x(4 * (i % 3)), G.w(4), CT + (i // 3) * (rh + gap)
        cv.line(f"Item rule {i + 1}", x, y, x + w, y, GREY)
        png, svg = a.icon(icon, GREEN)
        cv.picture(f"Icon {i + 1}", png, (x, y + 20, 40, 40), f"{head} icon", svg)
        cv.text(f"Item title {i + 1}", (x, y + 76, w, 22), P(head, 16, NAVY, HEAD, bold=True))
        cv.text(f"Item text {i + 1}", (x, y + 104, w, rh - 108), P(text, 12))


def s_image_text(cv: Canvas, a: Assets):
    title(cv, "Action title: one message, with a photo that shows its subject")
    w = G.w(6)
    cv.text("Subheading", (G.x(0), CT, w, 52),
            P("Use a photo that shows the place, the people or the asset in question", 18, NAVY, HEAD, bold=True,
              line=95))
    items = ["Choose from the BwB photo library: Photos in the BwB tab", "Crop the photo to its frame; never stretch it",
             "Keep text off the photo and credit it in the source line", "Add alt text that says what the photo shows"]
    cv.text("Text", (G.x(0), CT + 72, w, 180), "".join(P(t, 14, bullet="square", after=8) for t in items))
    callout(cv, (G.x(0), CB - 64, w, 64), "Key point",
            "What the reader should take from the photo and the text, in a line or two.")
    rid = a.image(photo("hero-solar-woman", w / CH, 0.68, 0.5, 1.0, 960))
    cv.picture("Photo", rid, (G.x(6), CT, w, CH), "A person standing beside a solar farm at dusk")
    source(cv, "Photo: solar farm at dusk, BwB image library. Credit every photo that is not BwB\u2019s own.")


CASE = [("The challenge", "Two or three sentences on the client\u2019s situation and the question they faced. Name the "
                          "constraint that made it hard."),
        ("What we did", "The approach in two or three sentences: the analysis, the structure we designed and the "
                        "partners we brought in.")]
RESULTS = [("USD 120m", "mobilised from private investors"), ("3x", "leverage on the public funds committed"),
           ("18", "months from mandate to financial close"), ("4", "countries covered by the programme")]


def s_case(cv: Canvas, a: Assets):
    title(cv, "Action title: what the case shows for the client\u2019s decision")
    px, pw, pad = G.x(0), G.w(4), 20
    cv.shape("Case panel", (px, CT, pw, CH), fill=NAVY)
    cv.text("Case label", (px + pad, CT + 24, pw - 2 * pad, 16), P("Case study", 12, LIME, bold=True))
    cv.text("Case title", (px + pad, CT + 48, pw - 2 * pad, 64),
            P("Project or client name, on one or two lines", 24, WHITE, HEAD, bold=True, line=95))
    cv.text("Case details", (px + pad, CT + 128, pw - 2 * pad, 80),
            P("Country, year", 14, WHITE, after=4) + P("Sector and instrument", 14, WHITE, after=4)
            + P("Client: sovereign, city or company", 14, WHITE))
    png, svg = a.icon("landmark", WHITE)
    cv.picture("Case icon", png, (px + pad, CB - pad - 40, 40, 40), "Landmark icon", svg)
    for i, (head, text) in enumerate(CASE):
        x, w = G.x(4 + 4 * i), G.w(4)
        cv.text(f"Case heading {i + 1}", (x, CT, w, 24), P(head, 16, NAVY, HEAD, bold=True))
        cv.text(f"Case text {i + 1}", (x, CT + 32, w, 120), P(text, 12))
    rule = CT + 164
    cv.line("Results rule", G.x(4), rule, RIGHT, rule, GREY)
    cv.text("Results heading", (G.x(4), rule + 14, G.w(8), 24), P("Results", 16, NAVY, HEAD, bold=True))
    for i, (value, label) in enumerate(RESULTS):
        x, w = G.x(4 + 2 * i), G.w(2)
        cv.text(f"Result {i + 1}", (x, rule + 48, w, 44), figure(value, 32, GREEN), anchor="b")
        cv.text(f"Result label {i + 1}", (x, rule + 100, w, 60), P(label, 12))
    source(cv, "Source: BwB project records (year). Use figures the client has agreed we can publish.")


def s_quote(cv: Canvas, a: Assets):
    title(cv, "Action title: what stakeholders told us, in their own words")
    cv.shape("Quote panel", (G.x(0), CT, G.w(12), CH), fill=NAVY)
    x, w = G.x(1), G.w(10)
    cv.text("Quote mark", (x, CT + 20, 80, 80), P("\u201C", 120, LIME, HEAD, bold=True, line=100))
    cv.text("Quote", (x, CT + 116, w, 130),
            P("Put here the one quotation that carries the point, in the speaker\u2019s own words, cut to two or "
              "three lines.", 28, WHITE, HEAD, italic=True, line=100))
    cv.text("Attribution", (x, CB - 92, w, 56),
            P("**Name Surname**, role, organisation", 14, WHITE, after=4)
            + P("Interview, October 2026. Quote with permission, or describe the speaker: \u2018a senior official at "
                "the ministry of finance\u2019.", 12, WHITE), anchor="b")


TEAM = [("team-discussion", (0.583, 0.242, 0.281), "Engagement lead",
         "Accountable for quality and for the client relationship. Give years of experience and two similar projects."),
        ("team-interview", (0.365, 0.461, 0.271), "Project manager",
         "Runs the workplan day to day. Name the sectors and the regions they know best."),
        ("team-meeting-table", (0.763, 0.466, 0.280), "Sector expert",
         "Brings the technical depth. Say what they have done that this project needs."),
        ("team-workshop", (0.659, 0.313, 0.313), "Financial analyst",
         "Builds the models and the evidence base. Keep every profile to the same length.")]


def s_team(cv: Canvas, a: Assets):
    title(cv, "Action title: the team brings the mix of skills this project needs")
    size = G.w(3)
    for i, (image, (cx, cy, scale), role, text) in enumerate(TEAM):
        x = G.x(3 * i)
        rid = a.image(photo(image, 1.0, cx, cy, scale, 480))
        cv.picture(f"Photo {i + 1}", rid, (x, CT, size, size), "Team member photo: replace with their portrait")
        cv.text(f"Name {i + 1}", (x, CT + size + 14, size, 22), P("Name Surname", 16, NAVY, HEAD, bold=True))
        cv.text(f"Role {i + 1}", (x, CT + size + 38, size, 18), P(role, 12, CHARCOAL))
        cv.text(f"Profile {i + 1}", (x, CT + size + 64, size, CB - CT - size - 64), P(text, 12))


def s_closing(cv: Canvas, a: Assets):
    cv.placeholder("Title", 'type="title"', PH("Thank you"))
    for i in (1, 2):
        cv.placeholder(f"Contact {i}", f'type="body" idx="{i}"',
                       PH("**Name Surname**") + PH("Role") + PH("name.surname@bwb.earth") + PH("Phone number"))


def s_disclaimer(cv: Canvas, a: Assets):
    title(cv, "Disclaimer")
    paras, note = read_disclaimer()
    half = (len(paras) + 1) // 2
    pad, height = 24, 264
    cv.shape("Disclaimer panel", (G.x(0), CT, G.w(12), height), fill=OFFWHITE, line=GREY20, line_w=0.5)
    col = (G.w(12) - 3 * pad) / 2
    for i, group in enumerate((paras[:half], paras[half:])):
        body = "".join(P(t, 11, after=8, line=105) for t in group)
        if i:
            body += P(note, 11, CHARCOAL, line=105)
        cv.text(f"Disclaimer text {i + 1}", (G.x(0) + pad + i * (col + pad), CT + pad, col, height - 2 * pad), body)


class Spec:
    def __init__(self, sid: int, slide_id: str, layout: str, build, name: str, category: str, description: str,
                 tags: list):
        self.sid, self.id, self.layout, self.build = sid, slide_id, layout, build
        self.title, self.category, self.description, self.tags = name, category, description, tags

    @property
    def slide_ref(self) -> str:
        return f"{self.sid}#{creation_id('slide/' + self.id)}"


# The sldId of each slide is fixed here, not by its position, so reordering the library keeps every slide's id.
SLIDES = [
    Spec(256, "title", "BwB Title", s_title, "Title", "Structure",
         "Opening slide on navy with the logo, the presentation title, a subtitle and the date.",
         ["title", "cover", "opening", "navy"]),
    Spec(257, "agenda", "BwB Title only", s_agenda, "Agenda", "Structure",
         "Five numbered sections, each with a one-line description.", ["agenda", "contents", "sections", "overview"]),
    Spec(258, "section-divider", "BwB Section divider", s_divider, "Section divider", "Structure",
         "Navy divider with a large lime section number, the section title and a line of context.",
         ["divider", "section", "chapter", "navy"]),
    Spec(259, "executive-summary", "BwB Title only", s_summary, "Executive summary", "Messages",
         "Three key messages, each with a green marker, a headline and the evidence behind it.",
         ["summary", "key messages", "findings", "answer first"]),
    Spec(260, "bullets", "BwB Title and content", s_bullets, "Text and callout", "Messages",
         "An action title, two levels of bullets and a callout for the key takeaway.",
         ["text", "bullets", "callout", "takeaway"]),
    Spec(261, "two-columns", "BwB Two content", s_two_columns, "Two columns", "Messages",
         "Two columns with headings and bullets, for options or periods side by side, and a recommendation.",
         ["comparison", "options", "before and after", "columns"]),
    Spec(262, "key-figures", "BwB Title only", s_key_figures, "Key figures", "Data",
         "Four headline numbers on a navy band, the first in lime, with a note under each.",
         ["figures", "statistics", "numbers", "kpi", "navy"]),
    Spec(263, "chart-and-text", "BwB Title only", s_chart, "Chart and text", "Data",
         "A native column chart in the BwB colours, ready for Edit Data, with a panel on what it shows.",
         ["chart", "graph", "data", "commentary"]),
    Spec(264, "table", "BwB Title only", s_table, "Table", "Data",
         "A data table with a navy header row, right-aligned figures and a total row.",
         ["table", "figures", "comparison", "data"]),
    Spec(265, "process", "BwB Title only", s_process, "Process", "Visuals",
         "Four chevrons for the steps of a process, with the timing, activities and output of each.",
         ["process", "steps", "approach", "methodology", "chevrons"]),
    Spec(266, "matrix", "BwB Title only", s_matrix, "Two-by-two matrix", "Visuals",
         "Four quadrants on impact and feasibility, with the options plotted and listed.",
         ["matrix", "prioritisation", "quadrants", "options"]),
    Spec(267, "timeline", "BwB Title only", s_timeline, "Timeline", "Visuals",
         "Six milestones on a line, grouped into phases, with the critical path below.",
         ["timeline", "milestones", "roadmap", "phases"]),
    Spec(268, "workplan", "BwB Title only", s_workplan, "Workplan", "Data",
         "A Gantt-style table of activities over twelve weeks, with deliverables marked.",
         ["workplan", "gantt", "schedule", "plan"]),
    Spec(269, "icon-grid", "BwB Title only", s_icons, "Icon grid", "Visuals",
         "Six items in a grid, each with a green Lucide icon, a title and a sentence.",
         ["icons", "capabilities", "services", "grid"]),
    Spec(270, "image-and-text", "BwB Title only", s_image_text, "Image and text", "Visuals",
         "A photo from the BwB library beside a heading, bullets and a key point.", ["photo", "image", "picture", "text"]),
    Spec(271, "case-study", "BwB Title only", s_case, "Case study", "Messages",
         "A navy panel naming the case, the challenge and the approach, and four results.",
         ["case study", "example", "project", "results", "credentials"]),
    Spec(272, "quote", "BwB Title only", s_quote, "Quote", "People",
         "One quotation in Candara italic on a navy panel, with the speaker\u2019s name and role.",
         ["quote", "testimonial", "stakeholder", "interview"]),
    Spec(273, "team", "BwB Title only", s_team, "Team", "People",
         "Four photo frames with names, roles and a short profile.", ["team", "people", "staff", "profiles"]),
    Spec(274, "closing", "BwB Closing", s_closing, "Closing", "Closing",
         "Navy closing slide with two contacts and the BwB website.", ["closing", "contact", "thank you", "end"]),
    Spec(275, "disclaimer", "BwB Title only", s_disclaimer, "Disclaimer", "Closing",
         "The standard BwB disclaimer and the FCA note, in an off-white panel.",
         ["disclaimer", "legal", "compliance", "fca"]),
]


def s_starter_title(cv: Canvas, a: Assets):
    for name, ph in (("Title", 'type="ctrTitle"'), ("Subtitle", 'type="subTitle" idx="1"'),
                     ("Date", 'type="body" sz="quarter" idx="13"')):
        cv.placeholder(name, ph)


def s_starter_content(cv: Canvas, a: Assets):
    cv.placeholder("Title", 'type="title"')
    cv.placeholder("Content", 'idx="1"')


# The starter deck's placeholders are empty, so PowerPoint shows the layouts' prompts, such as the action title prompt.
STARTER = [
    Spec(256, "starter-title", "BwB Title", s_starter_title, "Title", "Structure", "", []),
    Spec(257, "starter-content", "BwB Title and content", s_starter_content, "Title and content", "Messages", "", []),
]


# ---------------------------------------------------------------------------------------------------------------------
# Assembly, the index and the thumbnails

def slide_xml(shapes: str, key: str) -> bytes:
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sld {NS}><p:cSld>{sp_tree(shapes)}'
            f"{creation_ext('slide/' + key)}</p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>"
            ).encode("utf-8")


def fixed_xlsx(data: bytes) -> bytes:
    """The workbook behind the chart, with fixed dates and timestamps, so the build gives the same bytes."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            blob = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                blob = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", lambda m: m.group(1) + FIXED_ISO.encode(),
                              blob)
            zout.writestr(zipfile.ZipInfo(info.filename, FIXED_TIME), blob, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def deterministic(data: bytes) -> bytes:
    """Rewrites the package with fixed timestamps, pins the chart workbook's dates and declares SVG as a default type."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            blob = zin.read(info.filename)
            if info.filename == "[Content_Types].xml" and b"image/svg+xml" in blob:
                blob = re.sub(rb'<Override PartName="[^"]+\.svg" ContentType="image/svg\+xml"/>', b"", blob)
                blob = blob.replace(b'<Default Extension="xml"',
                                    b'<Default Extension="svg" ContentType="image/svg+xml"/><Default Extension="xml"')
            elif info.filename.endswith(".xlsx"):
                blob = fixed_xlsx(blob)
            zout.writestr(zipfile.ZipInfo(info.filename, FIXED_TIME), blob, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def build(specs: list, title: str) -> bytes:
    prs = Presentation(io.BytesIO(base_package(len(specs), title)))
    layouts = {layout.name: layout for layout in prs.slide_layouts}
    package = prs.part.package
    for n, spec in enumerate(specs, 1):
        slide = prs.slides.add_slide(layouts[spec.layout])
        cv = Canvas()
        spec.build(cv, Assets(slide.part, package))
        if spec.layout not in NAVY_LAYOUTS:
            cv.slide_number('type="sldNum" sz="quarter" idx="12"', str(n))
        slide.part._element = parse_xml(slide_xml(cv.xml(), spec.id))
    for sld_id, spec in zip(prs.slides._sldIdLst, specs):
        sld_id.set("id", str(spec.sid))
    buf = io.BytesIO()
    prs.save(buf)
    return deterministic(buf.getvalue())


def check_specs():
    ids = [s.id for s in SLIDES]
    sids = [s.sid for s in SLIDES]
    if len(set(ids)) != len(ids) or len(set(sids)) != len(sids):
        sys.exit("Slide ids and sldIds must be unique")
    for s in SLIDES + STARTER:
        if s.category not in CATEGORIES or s.layout not in LAYOUT_NAMES or not 256 <= s.sid < 2147483648:
            sys.exit(f"Slide {s.id}: unknown category or layout, or sldId out of range")


def write_index() -> Path:
    slides = [{"id": s.id, "title": s.title, "category": s.category, "description": s.description, "tags": s.tags,
               "slideId": s.slide_ref, "thumb": f"assets/slides/thumbs/{s.id}.png"} for s in SLIDES]
    index = {"version": 1, "file": PPTX_OUT.relative_to(ROOT / "public").as_posix(),
             "starterDeck": DECK_OUT.relative_to(ROOT / "public").as_posix(), "slides": slides}
    path = SLIDES_OUT / "index.json"
    path.write_text(json.dumps(index, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return path


def render(pptx: Path, preview: Path = None) -> int:
    """Renders every slide with LibreOffice, then writes the palette PNG thumbnails. Returns their total size."""
    for tool in ("soffice", "pdftoppm"):
        if not shutil.which(tool):
            sys.exit(f"{tool} is not installed; use --no-thumbs to skip the thumbnails")
    thumbs = SLIDES_OUT / "thumbs"
    thumbs.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="bwb_slides_") as tmp:
        tmp = Path(tmp)
        subprocess.run(["soffice", f"-env:UserInstallation={(tmp / 'profile').as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(tmp), str(pptx)], check=True, capture_output=True,
                       timeout=600)
        pdf = tmp / f"{pptx.stem}.pdf"
        if not pdf.exists():
            sys.exit("LibreOffice did not produce a PDF")
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(THUMB_WIDTH), "-scale-to-y", "-1", str(pdf),
                        str(tmp / "page")], check=True)
        pages = sorted(tmp.glob("page-*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[1]))
        if len(pages) != len(SLIDES):
            sys.exit(f"Rendered {len(pages)} pages for {len(SLIDES)} slides")
        keep = {f"{s.id}.png" for s in SLIDES}
        for old in thumbs.glob("*.png"):
            if old.name not in keep:
                old.unlink()
        total = 0
        for spec, page in zip(SLIDES, pages):
            with Image.open(page) as im:
                small = im.convert("RGB").quantize(colors=128, method=Image.Quantize.MEDIANCUT,
                                                   dither=Image.Dither.FLOYDSTEINBERG)
            path = thumbs / f"{spec.id}.png"
            small.save(path, optimize=True)
            total += path.stat().st_size
        if preview:
            preview.mkdir(parents=True, exist_ok=True)
            for old in preview.glob("slide-*.png"):
                old.unlink()
            subprocess.run(["pdftoppm", "-png", "-r", "144", str(pdf), str(preview / "slide")], check=True)
    return total


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-thumbs", action="store_true", help="skip the thumbnails (they need LibreOffice and pdftoppm)")
    ap.add_argument("--preview", type=Path, help="also write a 1920 px render of every slide to this folder")
    args = ap.parse_args()
    check_specs()
    PPTX_OUT.parent.mkdir(parents=True, exist_ok=True)
    SLIDES_OUT.mkdir(parents=True, exist_ok=True)
    PPTX_OUT.write_bytes(build(SLIDES, "BwB slide template and library"))
    DECK_OUT.write_bytes(build(STARTER, ""))
    index = write_index()
    print(f"{len(SLIDES)} slides, {PPTX_OUT.stat().st_size / 1e3:.0f} KB: {PPTX_OUT.relative_to(ROOT)}")
    print(f"{len(STARTER)} slides, {DECK_OUT.stat().st_size / 1e3:.0f} KB: {DECK_OUT.relative_to(ROOT)}")
    print(f"Index: {index.relative_to(ROOT)}")
    if not args.no_thumbs:
        total = render(PPTX_OUT, args.preview)
        print(f"Thumbnails: {total / 1e3:.0f} KB in {(SLIDES_OUT / 'thumbs').relative_to(ROOT)}")


if __name__ == "__main__":
    main()
