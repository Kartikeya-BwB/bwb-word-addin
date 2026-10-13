"""Assembly: builds a package from a list of slides, rewrites it with fixed timestamps so a rebuild gives the
same bytes, and writes the library index.
"""
import io
import json
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml import parse_xml

from .assets import Assets
from .catalogue import CATEGORIES, SLIDES
from .drawing import Canvas, NS, creation_ext, sp_tree
from .master import LAYOUT_NAMES, NAVY_LAYOUTS, base_package
from .sources import DECK_OUT, FIXED_ISO, FIXED_TIME, PPTX_OUT, ROOT, SLIDES_OUT
from .starter import STARTER


def slide_xml(shapes: str, key: str) -> bytes:
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sld {NS}><p:cSld>{sp_tree(shapes)}'
            f"{creation_ext('slide/' + key)}</p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>"
            ).encode("utf-8")


def fixed_xlsx(data: bytes) -> bytes:
    """The workbook behind the chart, with fixed dates and timestamps, so the build gives the same bytes."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            blob = zin.read(info.filename)
            if info.filename == "docProps/core.xml":
                blob = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", lambda m: m.group(1) + FIXED_ISO.encode(),
                              blob)
            zout.writestr(zipfile.ZipInfo(info.filename, FIXED_TIME), blob, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def deterministic(data: bytes) -> bytes:
    """Rewrites the package with fixed timestamps, pins the chart workbook's dates and declares SVG as a default type."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            blob = zin.read(info.filename)
            if info.filename == "[Content_Types].xml" and b"image/svg+xml" in blob:
                blob = re.sub(rb'<Override PartName="[^"]+\.svg" ContentType="image/svg\+xml"/>', b"", blob)
                blob = blob.replace(b'<Default Extension="xml"',
                                    b'<Default Extension="svg" ContentType="image/svg+xml"/><Default Extension="xml"')
            elif info.filename.endswith(".xlsx"):
                blob = fixed_xlsx(blob)
            zout.writestr(zipfile.ZipInfo(info.filename, FIXED_TIME), blob, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def build(specs: list, title: str) -> bytes:
    prs = Presentation(io.BytesIO(base_package(len(specs), title)))
    layouts = {layout.name: layout for layout in prs.slide_layouts}
    package = prs.part.package
    for n, spec in enumerate(specs, 1):
        slide = prs.slides.add_slide(layouts[spec.layout])
        cv = Canvas()
        spec.build(cv, Assets(slide.part, package))
        if spec.layout not in NAVY_LAYOUTS:
            cv.slide_number('type="sldNum" sz="quarter" idx="12"', str(n))
        slide.part._element = parse_xml(slide_xml(cv.xml(), spec.id))
    for sld_id, spec in zip(prs.slides._sldIdLst, specs):
        sld_id.set("id", str(spec.sid))
    buf = io.BytesIO()
    prs.save(buf)
    return deterministic(buf.getvalue())


def check_specs():
    ids = [s.id for s in SLIDES]
    sids = [s.sid for s in SLIDES]
    if len(set(ids)) != len(ids) or len(set(sids)) != len(sids):
        sys.exit("Slide ids and sldIds must be unique")
    for s in SLIDES + STARTER:
        if s.category not in CATEGORIES or s.layout not in LAYOUT_NAMES or not 256 <= s.sid < 2147483648:
            sys.exit(f"Slide {s.id}: unknown category or layout, or sldId out of range")


def write_index() -> Path:
    slides = [{"id": s.id, "title": s.title, "category": s.category, "description": s.description, "tags": s.tags,
               "slideId": s.slide_ref, "thumb": f"assets/slides/thumbs/{s.id}.png"} for s in SLIDES]
    index = {"version": 1, "file": PPTX_OUT.relative_to(ROOT / "public").as_posix(),
             "starterDeck": DECK_OUT.relative_to(ROOT / "public").as_posix(), "slides": slides}
    path = SLIDES_OUT / "index.json"
    path.write_text(json.dumps(index, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return path
